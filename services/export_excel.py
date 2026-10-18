import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from services import db

def generate_academic_workbook() -> io.BytesIO:
    """
    Generates a multi-sheet Microsoft Excel (.xlsx) workbook containing all 7 relational tables:
    Students, SchoolRecords, Semesters, StudentMarks, GapYears, Seminars, Activities.
    Returns a BytesIO stream ready for HTTP download.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    tables_config = [
        ('Students', "SELECT * FROM students ORDER BY id ASC"),
        ('SchoolRecords', "SELECT * FROM school_records ORDER BY id ASC"),
        ('Semesters', "SELECT * FROM semesters ORDER BY student_id ASC, semester_number ASC"),
        ('StudentMarks', "SELECT * FROM student_marks ORDER BY semester_id ASC, id ASC"),
        ('GapYears', "SELECT * FROM gap_years ORDER BY id ASC"),
        ('Seminars', "SELECT * FROM seminars ORDER BY event_date DESC"),
        ('Activities', "SELECT * FROM activities ORDER BY event_date DESC")
    ]

    # Styles
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')
    data_font = Font(name='Segoe UI', size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    alt_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    for sheet_name, query in tables_config:
        ws = wb.create_sheet(title=sheet_name)
        rows = db.query_all(query)

        if not rows:
            ws.append(["No records available"])
            continue

        headers = list(rows[0].keys())
        ws.append(headers)

        # Style header row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Add data rows
        for row_idx, r_dict in enumerate(rows, start=2):
            row_values = [r_dict[h] for h in headers]
            ws.append(row_values)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
